from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "rtsp_cctv_streaming_handoff.md"
OUT = ROOT / "rtsp_cctv_streaming_handoff.docx"


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run(run, size=11, bold=False, color="000000", font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill="F1F3F4"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)

    for name, size, color, before, after in [
        ("Heading 1", 20, "000000", 20, 6),
        ("Heading 2", 16, "000000", 18, 6),
        ("Heading 3", 14, "434343", 16, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=3)
    r = p.add_run(title)
    set_run(r, size=26, bold=False)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=14)
    r = p.add_run(subtitle)
    set_run(r, size=11, color="555555")


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=0, line=1.0)
        shade_paragraph(p)
        r = p.add_run(line if line else " ")
        set_run(r, size=9.5, font="Consolas", color="202124")
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=8)


def add_markdown_line(doc, line):
    stripped = line.strip()
    if not stripped:
        return
    if stripped.startswith("# "):
        add_title(doc, stripped[2:], "팀 공유용 RTSP CCTV 라이브 스트리밍 구축 기록")
        return
    if stripped.startswith("## "):
        doc.add_paragraph(stripped[3:], style="Heading 1")
        return
    if stripped.startswith("### "):
        doc.add_paragraph(stripped[4:], style="Heading 2")
        return
    if stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, before=0, after=4)
        r = p.add_run(stripped[2:])
        set_run(r)
        return
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    r = p.add_run(stripped)
    set_run(r)


def build():
    text = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    in_code = False
    code_lines = []
    for line in text:
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        add_markdown_line(doc, line)

    if code_lines:
        add_code_block(doc, code_lines)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
